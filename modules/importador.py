"""
modules/importador.py — El G POS
Importa inventario desde un Excel (hojas PRODUCTOS e INSUMOS) a la base local.

Diseñado para el documento 'Inventario - Import.xlsx':
  • Hoja PRODUCTOS: nombre, codigo, precio_venta, precio_costo, stock,
                    stock_minimo, categoria   (lo que se vende)
  • Hoja INSUMOS:   nombre, unidad, stock, stock_minimo   (ingredientes)

La importación es idempotente:
  - productos: hace UPSERT por 'codigo' (si el SKU ya existe, actualiza).
  - insumos:   hace UPSERT por 'nombre'.
Así se puede re-importar sin duplicar.
"""
from __future__ import annotations

from database import get_connection

# Categorías que se crean si faltan, con su tipo ('tienda' | 'cocina').
# Las demás categorías desconocidas se crean como 'tienda' por defecto.
_CATEGORIA_TIPO = {
    "Accesorios TCG": "tienda",
    "Juegos de mesa": "tienda",
    "TCG - Sobres": "tienda",
    "TCG - Mazos": "tienda",
    "TCG - Paquetes cartas": "tienda",
    "Mecato": "tienda",
    # Bebidas embotelladas = venta al detal → tipo tienda (aparecen en la
    # pestaña 'Tienda' de ventas directas junto al resto del mecato).
    "Bebidas": "tienda",
    "Comidas rápidas": "cocina",
    "Combos cocina": "cocina",
    # Platos de cocina (fichas técnicas)
    "Crepes salados": "cocina",
    "Crepes dulces": "cocina",
    "Sándwiches": "cocina",
    "Hamburguesas": "cocina",
    "Alitas": "cocina",
}

MARGEN_COSTO = 0.80  # precio_costo = precio_venta * 0.80 si viene vacío


def _leer_hoja(ws) -> list[dict]:
    """Lee una hoja como lista de dicts usando la fila 1 como encabezados."""
    headers = []
    for c in range(1, ws.max_column + 1):
        h = ws.cell(1, c).value
        headers.append(str(h).strip() if h is not None else None)
    filas = []
    for r in range(2, ws.max_row + 1):
        fila = {}
        vacia = True
        for c, h in enumerate(headers, 1):
            if not h:
                continue
            v = ws.cell(r, c).value
            if isinstance(v, str):
                v = v.strip()
            fila[h] = v
            if v not in (None, ""):
                vacia = False
        if not vacia:
            filas.append(fila)
    return filas


def _num(v, defecto=0.0):
    if v in (None, "", "-"):
        return defecto
    try:
        return float(v)
    except (TypeError, ValueError):
        return defecto


def _mapa_categorias(conn) -> dict[str, int]:
    return {n: i for i, n in conn.execute("SELECT id, nombre FROM categorias")}


def _asegurar_categoria(conn, nombre: str, cache: dict) -> int:
    nombre = (nombre or "").strip()
    if nombre in cache:
        return cache[nombre]
    tipo = _CATEGORIA_TIPO.get(nombre, "tienda")
    cur = conn.execute(
        "INSERT INTO categorias (nombre, tipo) VALUES (?, ?)", (nombre, tipo))
    cache[nombre] = cur.lastrowid
    return cache[nombre]


def importar_inventario(ruta_xlsx: str) -> dict:
    """
    Importa PRODUCTOS e INSUMOS desde el Excel. Retorna un resumen:
    {productos_nuevos, productos_actualizados, insumos_nuevos,
     insumos_actualizados, categorias_creadas, omitidos:[...], errores:[...]}
    """
    import openpyxl  # import perezoso (solo se usa al importar)

    wb = openpyxl.load_workbook(ruta_xlsx, data_only=True)
    res = {
        "productos_nuevos": 0, "productos_actualizados": 0,
        "insumos_nuevos": 0, "insumos_actualizados": 0,
        "categorias_creadas": 0, "categorias_ajustadas": 0,
        "recetas_lineas": 0, "preparaciones": 0, "recetas_prep": 0,
        "omitidos": [], "errores": [],
    }

    conn = get_connection()
    try:
        cat_cache = _mapa_categorias(conn)
        cat_iniciales = set(cat_cache)

        # Reconciliar el TIPO de las categorías conocidas cuya definición cambió
        # (p.ej. 'Bebidas' sembrada como cocina → debe ser tienda para que las
        # bebidas aparezcan en la pestaña 'Tienda' de ventas directas).
        for nombre, tipo_deseado in _CATEGORIA_TIPO.items():
            if nombre in cat_cache:
                row = conn.execute(
                    "SELECT tipo FROM categorias WHERE nombre = ?", (nombre,)).fetchone()
                if row and row[0] != tipo_deseado:
                    conn.execute(
                        "UPDATE categorias SET tipo = ? WHERE nombre = ?",
                        (tipo_deseado, nombre))
                    res["categorias_ajustadas"] += 1

        # ── PRODUCTOS ──────────────────────────────────────────────────────
        if "PRODUCTOS" in wb.sheetnames:
            for fila in _leer_hoja(wb["PRODUCTOS"]):
                nombre = (fila.get("nombre") or "").strip()
                if not nombre:
                    continue
                pv = fila.get("precio_venta")
                if pv in (None, "", "-"):
                    res["omitidos"].append(f"{nombre} (sin precio_venta)")
                    continue
                pv = _num(pv)
                pc = fila.get("precio_costo")
                pc = round(pv * MARGEN_COSTO) if pc in (None, "", "-") else _num(pc)
                codigo = fila.get("codigo") or None
                if isinstance(codigo, str):
                    codigo = codigo.strip() or None
                cat_id = _asegurar_categoria(conn, fila.get("categoria"), cat_cache)
                stock = _num(fila.get("stock"))
                sm = _num(fila.get("stock_minimo"))

                # ¿ya existe (por código)?
                existe = None
                if codigo:
                    row = conn.execute(
                        "SELECT id FROM productos WHERE codigo = ?", (codigo,)).fetchone()
                    existe = row[0] if row else None
                if existe:
                    conn.execute("""
                        UPDATE productos SET nombre=?, precio_venta=?, precio_costo=?,
                               stock=?, stock_minimo=?, categoria_id=?, activo=1
                        WHERE id=?
                    """, (nombre, pv, pc, stock, sm, cat_id, existe))
                    res["productos_actualizados"] += 1
                else:
                    conn.execute("""
                        INSERT INTO productos
                          (nombre, codigo, precio_venta, precio_costo,
                           stock, stock_minimo, categoria_id)
                        VALUES (?, ?, ?, ?, ?, ?, ?)
                    """, (nombre, codigo, pv, pc, stock, sm, cat_id))
                    res["productos_nuevos"] += 1

        # ── INSUMOS ────────────────────────────────────────────────────────
        if "INSUMOS" in wb.sheetnames:
            for fila in _leer_hoja(wb["INSUMOS"]):
                nombre = (fila.get("nombre") or "").strip()
                if not nombre:
                    continue
                unidad = (fila.get("unidad") or "unidad")
                if isinstance(unidad, str):
                    unidad = unidad.strip() or "unidad"
                stock = _num(fila.get("stock"))
                sm = _num(fila.get("stock_minimo"))
                row = conn.execute(
                    "SELECT id FROM insumos WHERE nombre = ?", (nombre,)).fetchone()
                if row:
                    conn.execute("""
                        UPDATE insumos SET stock=?, unidad=?, stock_minimo=?, activo=1
                        WHERE id=?
                    """, (stock, unidad, sm, row[0]))
                    res["insumos_actualizados"] += 1
                else:
                    conn.execute("""
                        INSERT INTO insumos (nombre, stock, unidad, stock_minimo)
                        VALUES (?, ?, ?, ?)
                    """, (nombre, stock, unidad, sm))
                    res["insumos_nuevos"] += 1

        # ── RECETAS ────────────────────────────────────────────────────────
        # Hoja opcional: columnas producto, insumo, cantidad. Vincula un plato
        # de cocina con sus insumos/preparaciones (para descontar al vender).
        # La receta de cada producto listado se REEMPLAZA (idempotente).
        if "RECETAS" in wb.sheetnames:
            prod_por_nombre = {n: i for i, n in
                               conn.execute("SELECT id, nombre FROM productos")}
            insumo_por_nombre = {n: i for i, n in
                                 conn.execute("SELECT id, nombre FROM insumos")}
            recetas = {}   # producto_id -> [(insumo_id, cantidad)]
            for fila in _leer_hoja(wb["RECETAS"]):
                pnombre = (fila.get("producto") or "").strip()
                inombre = (fila.get("insumo") or "").strip()
                cant = _num(fila.get("cantidad"))
                if not pnombre or not inombre or cant <= 0:
                    continue
                pid = prod_por_nombre.get(pnombre)
                iid = insumo_por_nombre.get(inombre)
                if not pid:
                    res["errores"].append(f"receta: producto '{pnombre}' no existe")
                    continue
                if not iid:
                    res["errores"].append(f"receta: insumo '{inombre}' no existe")
                    continue
                recetas.setdefault(pid, []).append((iid, cant))
            for pid, lineas in recetas.items():
                conn.execute("DELETE FROM receta_insumos WHERE producto_id = ?", (pid,))
                conn.executemany(
                    "INSERT INTO receta_insumos (producto_id, insumo_id, cantidad) "
                    "VALUES (?, ?, ?)", [(pid, iid, c) for iid, c in lineas])
                res["recetas_lineas"] += len(lineas)

        # ── PREPARACIONES ──────────────────────────────────────────────────
        # Hoja opcional: columnas preparacion, rendimiento, componente, cantidad.
        # Marca un insumo como preparación (con su rendimiento por lote) y define
        # su receta de producción. Un componente puede ser un crudo u otra
        # preparación (anidado). La receta de cada preparación se REEMPLAZA.
        if "PREPARACIONES" in wb.sheetnames:
            insumo_por_nombre = {n: i for i, n in
                                 conn.execute("SELECT id, nombre FROM insumos")}
            rendimientos = {}   # preparacion_id -> rendimiento
            preps = {}          # preparacion_id -> [(insumo_id, cantidad)]
            for fila in _leer_hoja(wb["PREPARACIONES"]):
                pnombre = (fila.get("preparacion") or "").strip()
                cnombre = (fila.get("componente") or "").strip()
                cant = _num(fila.get("cantidad"))
                rend = _num(fila.get("rendimiento"))
                if not pnombre:
                    continue
                pid = insumo_por_nombre.get(pnombre)
                if not pid:
                    res["errores"].append(f"preparación: '{pnombre}' no existe como insumo")
                    continue
                if rend > 0:
                    rendimientos[pid] = rend
                if not cnombre or cant <= 0:
                    continue
                cid = insumo_por_nombre.get(cnombre)
                if not cid:
                    res["errores"].append(f"preparación '{pnombre}': componente '{cnombre}' no existe")
                    continue
                preps.setdefault(pid, []).append((cid, cant))
            for pid in set(rendimientos) | set(preps):
                rend = rendimientos.get(pid, 0)
                conn.execute("""
                    INSERT INTO preparaciones (insumo_id, rendimiento) VALUES (?, ?)
                    ON CONFLICT(insumo_id) DO UPDATE SET rendimiento = excluded.rendimiento
                """, (pid, rend))
                lineas = preps.get(pid, [])
                conn.execute("DELETE FROM receta_preparacion WHERE preparacion_id = ?", (pid,))
                conn.executemany(
                    "INSERT INTO receta_preparacion (preparacion_id, insumo_id, cantidad) "
                    "VALUES (?, ?, ?)", [(pid, cid, c) for cid, c in lineas])
                res["preparaciones"] += 1
                res["recetas_prep"] += len(lineas)

        conn.commit()
        res["categorias_creadas"] = len(set(cat_cache) - cat_iniciales)
    except Exception as e:
        conn.rollback()
        res["errores"].append(str(e))
        raise
    finally:
        conn.close()

    return res
